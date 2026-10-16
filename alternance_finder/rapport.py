"""Exports : classeur Excel + tableau de bord HTML interactif."""
import html
import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .cv import ProfilCV
from .models import Offre

ENTETE = PatternFill("solid", fgColor="1F4E78")


def _emails(o: Offre, sep: str = "\n") -> str:
    return sep.join(
        f"{e['email']}" + (f" ({e['nom']})" if e.get("nom") else "") + f" [{e['source']}]" for e in o.emails_rh
    )


def _feuille(wb, titre: str, colonnes: list[tuple[str, int]], lignes: list[list], col_score: int | None = None):
    ws = wb.create_sheet(titre)
    ws.append([c for c, _ in colonnes])
    for i, (_, largeur) in enumerate(colonnes, 1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True, color="FFFFFF"), ENTETE
        ws.column_dimensions[get_column_letter(i)].width = largeur
    for ligne in lignes:
        ws.append(ligne)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(cell.value, str) and cell.value.startswith("http"):
                cell.hyperlink, cell.font = cell.value, Font(color="0563C1", underline="single")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if col_score and ws.max_row > 1:
        lettre = get_column_letter(col_score)
        ws.conditional_formatting.add(
            f"{lettre}2:{lettre}{ws.max_row}",
            ColorScaleRule(start_type="num", start_value=0, start_color="F8696B",
                           mid_type="num", mid_value=50, mid_color="FFEB84",
                           end_type="num", end_value=100, end_color="63BE7B"))
    return ws


def excel(chemin: Path, profil: ProfilCV, offres: list[Offre], spontanees: list[Offre], a_travailler):
    wb = Workbook()
    wb.remove(wb.active)

    _feuille(wb, "Offres matchées", [
        ("Score", 8), ("Nouveau", 9), ("Catégorie", 16), ("Poste", 40), ("Entreprise", 28), ("Lieu", 22),
        ("Contrat", 18), ("Date", 11), ("Source", 18), ("Compétences OK", 35), ("Compétences manquantes", 35),
        ("E-mails RH", 45), ("Téléphone", 15), ("Lien offre", 45), ("Recherche RH LinkedIn", 30), ("Site web", 30),
    ], [[
        o.score, "oui" if o.nouveau else "", o.categorie, o.titre, o.entreprise, o.lieu, o.type_contrat, o.date,
        o.source, ", ".join(o.competences_ok), ", ".join(o.competences_manquantes), _emails(o), o.telephone, o.url,
        o.lien_recherche_rh, o.site_web,
    ] for o in offres], col_score=1)

    _feuille(wb, "Candidatures spontanées", [
        ("Pertinence", 11), ("Nouveau", 9), ("Entreprise", 32), ("Secteur (NAF)", 38), ("Adresse", 35),
        ("Site web", 30), ("Page recrutement", 40), ("E-mails RH", 45), ("Téléphone", 15),
        ("Postuler (La Bonne Alternance)", 45), ("Recherche RH LinkedIn", 30), ("Chercher le site", 30), ("SIRET", 16),
    ], [[
        o.score, "oui" if o.nouveau else "", o.entreprise, o.secteur, o.lieu, o.site_web, o.page_recrutement,
        _emails(o), o.telephone, o.url, o.lien_recherche_rh, o.lien_recherche_site, o.siret,
    ] for o in spontanees], col_score=1)

    contacts = []
    for o in offres + spontanees:
        for e in o.emails_rh:
            contacts.append([o.entreprise, e["email"], e.get("nom", ""), e.get("poste", ""), e["source"],
                             o.titre, o.url])
    _feuille(wb, "Contacts RH", [("Entreprise", 30), ("E-mail", 38), ("Nom", 22), ("Fonction", 25),
                                 ("Source", 18), ("Pour", 35), ("Lien", 45)], contacts)

    ws = _feuille(wb, "Mon profil", [("Élément", 30), ("Valeur", 90)], [
        ["CV", profil.fichier], ["E-mail", profil.email], ["Téléphone", profil.telephone],
        ["Postes visés", profil.cible], ["Compétences détectées", ", ".join(sorted(profil.competences))],
        ["", ""], ["Compétences les plus demandées absentes du CV", ""],
        *[[c, f"demandée dans {n} offre(s)"] for c, n in a_travailler],
    ])
    ws.auto_filter.ref = None
    wb.save(chemin)


def html_dashboard(chemin: Path, profil: ProfilCV, offres: list[Offre], spontanees: list[Offre], a_travailler):
    data = [o.to_dict() for o in offres + spontanees]
    for d in data:
        d["description"] = d["description"][:1500]
    page = (Path(__file__).parent / "dashboard.html").read_text(encoding="utf-8")
    page = (page
            .replace("/*DATA*/[]", json.dumps(data, ensure_ascii=False).replace("</", "<\\/"))
            .replace("/*PROFIL*/{}", json.dumps({
                "competences": sorted(profil.competences), "cible": profil.cible,
                "a_travailler": a_travailler, "fichier": profil.fichier,
            }, ensure_ascii=False).replace("</", "<\\/"))
            .replace("{{DATE}}", html.escape(datetime.now().strftime("%d/%m/%Y %H:%M"))))
    chemin.write_text(page, encoding="utf-8")
